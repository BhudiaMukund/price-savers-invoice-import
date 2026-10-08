"""Builds the sample catalogues used by the tests and demo mode (python3 test/fixtures/make.py).
Suppliers are fictional and match the sample invoices: Harbour Novelty Imports, Sunburst Party Supplies, Koala Craft."""
import csv, io, os, re, zipfile
from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
from openpyxl.styles import Font

OUT = os.path.dirname(os.path.abspath(__file__))
FONT = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"

def ean(body):
    s = sum(int(c) * (3 if i % 2 == 0 else 1) for i, c in enumerate(reversed(body)))
    return body + str((10 - s % 10) % 10)

PALETTE = ["#F4A261", "#2A9D8F", "#E76F51", "#8AB17D", "#7B5EA7", "#E9C46A", "#457B9D", "#D62828"]

def photo(label, sub, i, fmt="JPEG", size=480):
    im = Image.new("RGB", (size, size), "#FAFAF7")
    d = ImageDraw.Draw(im)
    col = PALETTE[i % len(PALETTE)]
    k = size / 480
    d.rounded_rectangle((70 * k, 60 * k, size - 70 * k, size - 150 * k), radius=40 * k, fill=col)
    d.ellipse((size / 2 - 60 * k, 120 * k, size / 2 + 60 * k, 240 * k), fill="#FFFFFF")
    f = ImageFont.truetype(FONT, max(12, int(34 * k)))
    g = ImageFont.truetype(FONT, max(10, int(24 * k)))
    d.text((size / 2, size - 110 * k), label, font=f, fill="#2B2F39", anchor="mm")
    d.text((size / 2, size - 70 * k), sub, font=g, fill="#6B7280", anchor="mm")
    b = io.BytesIO()
    im.save(b, fmt, quality=82)
    return b.getvalue()

# ------------------------------------------------------------ 1. Excel with pictures over the cells
HN = [  # code, barcode (int = stored as a number), description, photos
    ("HN-7701", int(ean("939911177013")), "SKULL CANDLE HOLDER BLK 12CM", ["front", "back"]),
    ("HN-7705", int(ean("939911177051")), "SPIDER WEB STRETCH W/ 2 SPIDERS 60G", ["front"]),
    ("HN-7712", int(ean("939911177129")), "PUMPKIN BUCKET LED 18CM ORANGE", ["front"]),
    ("HN-8001", int(ean("07123488001")), "GLOW STICKS 20CM 15PK TUBE", ["front"]),
    ("HN-8003", int(ean("07123488003")), "GLOW BRACELETS 50PK ASSORTED COLOURS", ["pink", "green", "blue"]),
    ("HN-8010", None, "LED FINGER LIGHTS 4PC", ["front"]),
    ("HN-6120", 6921866190107, "AUSSIE FLAG HEADBAND", ["front"]),
    ("HN-6135", int(ean("939911161350")), "NOVELTY GLASSES RED", ["front", "side"]),
    ("HN-9100", int(ean("939911191002")), "WITCH HAT BLACK ADULT", ["front"]),
    ("HN-9105", int(ean("939911191057")), "VAMPIRE TEETH GLOW 2PK", []),
]

def build_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Range 2026"
    ws["A1"] = "HARBOUR NOVELTY IMPORTS - PRODUCT RANGE 2026"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Sample catalogue for Price Savers staff training. Supplier is fictional."
    for i, h in enumerate(["Item No", "Barcode", "Description", "Inner", "RRP", "Picture"], 1):
        ws.cell(4, i, h).font = Font(bold=True)
    r = 5
    k = 0
    for code, bc, desc, pics in HN:
        ws.cell(r, 1, code); ws.cell(r, 2, bc); ws.cell(r, 3, desc); ws.cell(r, 4, 6); ws.cell(r, 5, 4.95)
        ws.row_dimensions[r].height = 60
        for j, p in enumerate(pics):
            img = XLImage(io.BytesIO(photo(code, p, k)))
            img.width = img.height = 70
            col = 6 + j
            if code == "HN-8001":
                # Spans two rows, starting in the row above: should still land on HN-8001's row.
                a = TwoCellAnchor()
                a._from = AnchorMarker(col=col - 1, row=r - 2, rowOff=600000)  # near the bottom of the row above
                a.to = AnchorMarker(col=col, row=r - 1, rowOff=500000)  # most of it on HN-8001's row
                img.anchor = a
            else:
                img.anchor = f"{chr(64 + col)}{r}"
            ws.add_image(img)
            k += 1
        r += 1
    ws.cell(r + 1, 1, "All prices ex GST. E&OE.")
    # A logo in the heading area, not on any product row
    logo = XLImage(io.BytesIO(photo("HNI", "logo", 7, size=200)))
    logo.width = logo.height = 40
    logo.anchor = "F1"
    ws.add_image(logo)
    for col, w in zip("ABCDEF", [11, 16, 40, 7, 8, 12]):
        ws.column_dimensions[col].width = w
    path = os.path.join(OUT, "harbour-catalogue.xlsx")
    wb.save(path)
    return path

# ------------------------------------------------------------ 2. Excel with pictures placed in cells (Excel 365)
def build_incell():
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    rows = [("Code", "EAN", "Product", "Photo"), ("KC-118", ean("939902211806"), "Glitter Glue Pens 6pk", "IMG1"), ("KC-512", ean("939902251207"), "Craft Pom Poms 100pc", "IMG2")]
    for r, vals in enumerate(rows, 1):
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v)
    tmp = os.path.join(OUT, "_incell.xlsx")
    wb.save(tmp)
    zin = zipfile.ZipFile(tmp)
    parts = {n: zin.read(n) for n in zin.namelist()}
    zin.close()
    os.remove(tmp)
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    for n, ref in [(1, "D2"), (2, "D3")]:
        sheet = re.sub(rf'<c r="{ref}"[^>]*>.*?</c>', f'<c r="{ref}" t="e" vm="{n}"><v>#VALUE!</v></c>', sheet)
    parts["xl/worksheets/sheet1.xml"] = sheet.encode()
    parts["xl/media/image1.png"] = photo("KC-118", "in cell", 1, "PNG", 200)
    parts["xl/media/image2.png"] = photo("KC-512", "in cell", 2, "PNG", 200)
    R = "http://schemas.openxmlformats.org/package/2006/relationships"
    parts["xl/richData/richValueRel.xml"] = (
        '<richValueRels xmlns="http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><rel r:id="rId1"/><rel r:id="rId2"/></richValueRels>'
    ).encode()
    parts["xl/richData/_rels/richValueRel.xml.rels"] = (
        f'<Relationships xmlns="{R}">'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image2.png"/>'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/>'
        "</Relationships>"
    ).encode()
    # Values listed in the opposite order to the pictures, to prove the indexes are followed.
    parts["xl/richData/rdrichvalue.xml"] = (
        '<rvData xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata" count="2">'
        '<rv s="0"><v>1</v><v>5</v></rv><rv s="0"><v>0</v><v>5</v></rv></rvData>'
    ).encode()
    parts["xl/richData/rdrichvaluestructure.xml"] = (
        '<rvStructures xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata" count="1">'
        '<s t="_localImage"><k n="_rvRel:LocalImageIdentifier" t="i"/><k n="CalcOrigin" t="i"/></s></rvStructures>'
    ).encode()
    fut = lambda i: f'<bk><extLst><ext uri="{{3e2802c4-a4d2-4d8b-9148-e3be6c30e623}}"><xlrd:rvb i="{i}"/></ext></extLst></bk>'
    parts["xl/metadata.xml"] = (
        '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:xlrd="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata">'
        '<metadataTypes count="1"><metadataType name="XLRICHVALUE" minSupportedVersion="120000" copy="1" pasteAll="1" pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" clearComments="1" assign="1" coerce="1"/></metadataTypes>'
        f'<futureMetadata name="XLRICHVALUE" count="2">{fut(1)}{fut(0)}</futureMetadata>'
        '<valueMetadata count="2"><bk><rc t="1" v="0"/></bk><bk><rc t="1" v="1"/></bk></valueMetadata></metadata>'
    ).encode()
    rels = parts["xl/_rels/workbook.xml.rels"].decode()
    extra = (
        '<Relationship Id="rIdM" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sheetMetadata" Target="metadata.xml"/>'
        '<Relationship Id="rIdR1" Type="http://schemas.microsoft.com/office/2022/10/relationships/richValueRel" Target="richData/richValueRel.xml"/>'
        '<Relationship Id="rIdR2" Type="http://schemas.microsoft.com/office/2017/06/relationships/rdRichValue" Target="richData/rdrichvalue.xml"/>'
        '<Relationship Id="rIdR3" Type="http://schemas.microsoft.com/office/2017/06/relationships/rdRichValueStructure" Target="richData/rdrichvaluestructure.xml"/>'
    )
    parts["xl/_rels/workbook.xml.rels"] = rels.replace("</Relationships>", extra + "</Relationships>").encode()
    ct = parts["[Content_Types].xml"].decode()
    ct = ct.replace(
        "</Types>",
        '<Default Extension="png" ContentType="image/png"/>'
        '<Override PartName="/xl/metadata.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/>'
        '<Override PartName="/xl/richData/richValueRel.xml" ContentType="application/vnd.ms-excel.richvaluerel+xml"/>'
        '<Override PartName="/xl/richData/rdrichvalue.xml" ContentType="application/vnd.ms-excel.rdrichvalue+xml"/>'
        '<Override PartName="/xl/richData/rdrichvaluestructure.xml" ContentType="application/vnd.ms-excel.rdrichvaluestructure+xml"/>'
        "</Types>",
    )
    parts["[Content_Types].xml"] = ct.encode()
    path = os.path.join(OUT, "koala-incell.xlsx")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)

# ------------------------------------------------------------ 3. CSV with photo links (some "expired")
def build_csv():
    rows = [
        ("SB-1021", ean("939900110021"), "Balloon Latex 30cm Metallic Gold 25pk", "https://images.sunburst.example/SB-1021.jpg", "https://images.sunburst.example/SB-1021-b.jpg"),
        ("SB-1022", ean("939900110038"), "Balloon Latex 30cm Metallic Silver 25pk", "https://images.sunburst.example/SB-1022.jpg", ""),
        ("SB-2210", ean("939900122109"), "Happy Birthday Banner Holographic 2.1m", "https://images.sunburst.example/SB-2210.jpg?expires=1700000000&sig=expired", ""),
        ("SB-4410", ean("939900144101"), "Plates Paper 23cm Round White 20pc", "https://images.sunburst.example/SB-4410.jpg", ""),
        ("SB-4411", ean("939900144118"), "Cups Paper 260ml White 25pc", "https://images.sunburst.example/SB-4411.jpg", ""),
        ("SB-5120", "", "Candle Numeral #5 Glitter Pink", "https://images.sunburst.example/SB-5120.jpg", ""),
        ("SB-6001", ean("939900160019"), "Napkins Lunch 2ply Black 20pk", "", ""),
        ("SB-1031", ean("939900110311"), "Balloon Latex 30cm Rose Gold 25pk", "https://images.sunburst.example/SB-1031.jpg", ""),
    ]
    with open(os.path.join(OUT, "sunburst-links.csv"), "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["Stock Code", "EAN", "Description", "Image URL 1", "Image URL 2", "Pack Qty"])
        for r in rows:
            w.writerow([*r, 12])

# ------------------------------------------------------------ 4. ZIP of photos named by code or barcode
def build_zip():
    files = {
        "KC-118.jpg": ("KC-118", "front"),
        "KC-118_2.jpg": ("KC-118", "back"),
        f"{ean('939902220501')}.jpg": ("KC-205", "by barcode"),
        "KC-310-1.jpg": ("KC-310", "1"),
        "KC-310-2.jpg": ("KC-310", "2"),
        "KC-402 front.jpg": ("KC-402", "front"),
        "KC-520.png": ("KC-520", "png"),
    }
    with zipfile.ZipFile(os.path.join(OUT, "koala-photos.zip"), "w") as z:
        for i, (name, (code, sub)) in enumerate(files.items()):
            z.writestr(f"Koala photos/{name}", photo(code, sub, i, "PNG" if name.endswith(".png") else "JPEG", 360))
        z.writestr("__MACOSX/Koala photos/._KC-118.jpg", b"junk")
        z.writestr("Koala photos/Thumbs.db", b"junk")
        z.writestr("Koala photos/logo.emf", b"emf")

build_xlsx(); build_incell(); build_csv(); build_zip()
print(sorted(os.listdir(OUT)))
